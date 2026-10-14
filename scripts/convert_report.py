#!/usr/bin/env python3
"""Convert the project's HTML-based report-template.md to editable DOCX.

Only python-docx is required. This is a scoped converter for semantic HTML
embedded in Markdown, not a general Markdown/CSS rendering engine. Unsupported
markup raises an error instead of silently removing report content.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError:
    sys.exit('Нужен пакет python-docx в выбранной среде Python. Автоматическая установка не выполняется.')


@dataclass
class Element:
    tag: str
    attrs: dict = field(default_factory=dict)
    children: list = field(default_factory=list)

    def text(self):
        return ''.join(c if isinstance(c, str) else c.text() for c in self.children)


class ReportParser(HTMLParser):
    allowed = {'div','h1','h2','h3','p','table','thead','tbody','tr','td','th',
               'br','strong','b','em','i','u','a','span','img','ul','ol','li'}
    void = {'br','img'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Element('root')
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        if tag not in self.allowed:
            raise ValueError(f'Неподдерживаемый HTML-тег <{tag}>. Конвертация остановлена.')
        node = Element(tag, dict(attrs))
        if any(k.lower().startswith('on') for k in node.attrs):
            raise ValueError('Обработчики событий в отчёте не поддерживаются.')
        self.stack[-1].children.append(node)
        if tag not in self.void:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self.void:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if tag in self.void:
            return
        if len(self.stack) < 2 or self.stack[-1].tag != tag:
            raise ValueError(f'Некорректное закрытие </{tag}>.')
        self.stack.pop()

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def descendants(node, tag):
    for c in node.children:
        if isinstance(c, Element):
            if c.tag == tag:
                yield c
            yield from descendants(c, tag)


def css(node):
    return dict((k.strip().lower(), v.strip().lower()) for k,v in
                (part.split(':',1) for part in node.attrs.get('style','').split(';') if ':' in part))


def format_font(font, size):
    font.name = 'Times New Roman'
    font.size = Pt(size)
    font.color.rgb = RGBColor(0,0,0)


def format_run(run, size):
    format_font(run.font, size)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','eastAsia','cs']:
        rfonts.set(qn('w:'+key), 'Times New Roman')


def configure(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
    for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
        s=doc.styles[name]
        format_font(s.font, 16 if name=='Heading 1' else 12 if name=='Caption' else 14)
        for key in ['ascii','hAnsi','eastAsia','cs']:
            s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+key),'Times New Roman')
        pf=s.paragraph_format
        pf.line_spacing = 1 if name=='Caption' else 1.5
        pf.space_before = Pt(8) if name.startswith('Heading') else Pt(0)
        pf.space_after = Pt(6)
        pf.first_line_indent = Cm(1.25) if name=='Normal' else Cm(0)
        pf.widow_control = True
        pf.keep_with_next = name.startswith('Heading')
        pf.keep_together = name.startswith('Heading') or name=='Caption'
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if name in ['Heading 1','Caption'] else WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = 'Шаблон отчёта по учебной практике'


class Converter:
    def __init__(self, source):
        self.source = source
        self.doc = Document()
        configure(self.doc)
        self.figures = []
        self.frames = 0

    def inline(self, paragraph, children, size=14, bold=False, italic=False, underline=False):
        for child in children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                format_run(run,size)
                run.bold = True if bold else None
                run.italic = True if italic else None
                run.underline = True if underline else None
                continue
            tag=child.tag
            if tag=='br':
                paragraph.add_run().add_break()
            elif tag=='img':
                raw=child.attrs.get('src','')
                if not raw or '://' in raw or raw.startswith(('data:','/','\\')):
                    raise ValueError('Картинки должны иметь относительный локальный путь внутри папки отчёта.')
                path=(self.source.parent / raw).resolve()
                if not path.is_relative_to(self.source.parent.resolve()) or not path.is_file():
                    raise ValueError(f'Недоступно локальное изображение: {raw}')
                paragraph.add_run().add_picture(str(path),width=Cm(15.7))
            elif tag=='a':
                target=child.attrs.get('href','')
                if not target.lower().startswith(('https://','http://','mailto:')):
                    raise ValueError('Ссылки в Word должны иметь полный HTTPS/HTTP/mailto адрес.')
                rid=paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True)
                link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rid)
                run=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
                for name,value in [('rFonts',None),('sz',str(size*2)),('color','000000'),('u','single')]:
                    el=OxmlElement('w:'+name)
                    if name=='rFonts':
                        for k in ['ascii','hAnsi','eastAsia','cs']:el.set(qn('w:'+k),'Times New Roman')
                    else:el.set(qn('w:val'),value)
                    pr.append(el)
                run.append(pr); t=OxmlElement('w:t'); t.text=child.text(); run.append(t); link.append(run)
                paragraph._p.append(link)
            elif tag in {'strong','b','em','i','u','span'}:
                self.inline(paragraph,child.children,size,bold or tag in {'strong','b'},italic or tag in {'em','i'},underline or tag=='u')
            else:
                raise ValueError(f'Тег <{tag}> недопустим внутри абзаца.')

    def paragraph(self,node,parent=None,table=False):
        parent=parent or self.doc
        text=node.text().strip()
        if node.tag.startswith('h'):
            level=int(node.tag[1]); p=parent.add_paragraph(style=f'Heading {level}')
            size=16 if level==1 else 14
        elif re.match(r'^Рисунок\s+\d+\s+[—-]',text):
            p=parent.add_paragraph(style='Caption');size=12
            self.figures.append(int(re.match(r'^Рисунок\s+(\d+)',text).group(1)))
        else:
            p=parent.add_paragraph();size=12 if table else 14
            if table:
                p.paragraph_format.line_spacing=1
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.space_after=Pt(3)
        if css(node).get('page-break-before')=='always':p.paragraph_format.page_break_before=True
        self.inline(p,node.children,size)
        return p

    def table(self,node):
        rows=list(descendants(node,'tr'))
        cells=[[c for c in r.children if isinstance(c,Element) and c.tag in {'td','th'}] for r in rows]
        if not cells or not cells[0] or any(len(row)!=len(cells[0]) for row in cells):
            raise ValueError('Таблица должна иметь непустые строки с одинаковым числом ячеек.')
        if any('rowspan' in c.attrs or 'colspan' in c.attrs for row in cells for c in row):
            raise ValueError('Объединённые ячейки не поддерживаются; сохраните простую таблицу.')
        frame=len(cells)==1 and len(cells[0])==1 and (
            'Вставьте скриншот' in cells[0][0].text()
            or float(cells[0][0].attrs.get('height','0'))>=220)
        table=self.doc.add_table(rows=len(cells),cols=len(cells[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        width=Cm(16.5)
        pr=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            edge=OxmlElement('w:'+side)
            for k,v in [('val','single'),('sz','8'),('color','000000')]:edge.set(qn('w:'+k),v)
            borders.append(edge)
        pr.append(borders)
        for row,htmlcells in zip(table.rows,cells):
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for cell,htmlcell in zip(row.cells,htmlcells):
                cell.width=int(width/len(htmlcells))
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                # Delete the empty default paragraph only after creating actual content.
                default=cell.paragraphs[0]
                blocks=[c for c in htmlcell.children if isinstance(c,Element) and c.tag in {'p','h1','h2','h3'}]
                if blocks:
                    for c in htmlcell.children:
                        if isinstance(c,str) and c.strip():raise ValueError('Текст ячейки должен находиться внутри <p>.')
                        if isinstance(c,Element):self.paragraph(c,cell,table=True)
                else:
                    self.inline(default,htmlcell.children,12)
                if blocks:default._element.getparent().remove(default._element)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Cm(0)
                    if frame:
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.keep_with_next=True
                        p.paragraph_format.keep_together=True
        if frame:
            self.frames+=1
            px=float(cells[0][0].attrs.get('height','220'))
            table.rows[0].height=Pt(max(220,px)*0.75)
            table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST

    def blocks(self,node):
        for child in node.children:
            if isinstance(child,str):
                if child.strip():
                    raise ValueError('Вне HTML найден текст Markdown. Используйте семантические <h1>/<h2>/<h3>/<p>, как в report-template.md.')
                continue
            if child.tag=='div':self.blocks(child)
            elif child.tag in {'h1','h2','h3','p'}:self.paragraph(child)
            elif child.tag=='table':self.table(child)
            elif child.tag in {'ol','ul'}:
                for li in child.children:
                    if isinstance(li,str) and not li.strip():continue
                    if not isinstance(li,Element) or li.tag!='li':raise ValueError('Некорректный список.')
                    p=self.doc.add_paragraph(style='List Number' if child.tag=='ol' else 'List Bullet')
                    self.inline(p,li.children)
            else:raise ValueError(f'Неподдерживаемый блок <{child.tag}>.')

    def convert(self,output):
        parser=ReportParser();parser.feed(self.source.read_text(encoding='utf-8-sig'));parser.close()
        if len(parser.stack)!=1:raise ValueError('В HTML остались незакрытые теги.')
        self.blocks(parser.root)
        if self.frames!=len(self.figures):
            raise ValueError(f'Число рамок ({self.frames}) и подписей ({len(self.figures)}) не совпадает.')
        if self.figures!=list(range(1,len(self.figures)+1)):
            raise ValueError('Подписи рисунков должны быть пронумерованы последовательно с 1.')
        if output.exists():raise ValueError('Выходной файл существует. Выберите новое имя; существующие документы не перезаписываются.')
        output.parent.mkdir(parents=True,exist_ok=True)
        self.doc.save(output)
        # Reopen the actual deliverable before declaring success.
        saved=Document(output)
        if len(saved.tables)!=len(self.doc.tables):raise ValueError('Не все таблицы сохранены.')
        return self.frames


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('input',type=Path);p.add_argument('output',type=Path)
    args=p.parse_args()
    try:
        if args.input.suffix.lower()!='.md' or args.output.suffix.lower()!='.docx':
            raise ValueError('Ожидаются входной .md и выходной .docx.')
        frames=Converter(args.input.resolve()).convert(args.output.resolve())
        print(f'Word сохранён: {args.output.name}; рамок и подписей: {frames}. Требуется визуальная проверка в Word или через отрисовку.')
    except (ValueError,OSError) as exc:
        p.exit(1,f'Ошибка: {exc}\n')


if __name__=='__main__':main()
